#!/usr/bin/env python3
"""Merge the tester's spreadsheet verdicts into the collected run reports.

    python deploy/merge_tester_xlsx.py <file.xlsx> [reports-dir] [--sheet Прогоны] [--dry-run]

The verdicts the tester typed into the app are a short form. The real ones
live in their spreadsheet — longer, and covering runs that were never given
feedback in the app at all. This reads the sheet, matches rows to runs by
`request_id`, and writes the spreadsheet text into each report's `feedback`.

Columns are detected by content rather than by name, because a spreadsheet
someone maintains by hand gets its headings renamed and its columns reordered,
and a script that hard-codes "column D" silently merges the wrong field after
the first edit:

* the id column is the one whose cells look like `grade_<hex>`;
* the verdict column is the one saying прав/не прав/верно/неверно;
* the comment column is the longest free text that is not the id;
* expected scores are read only from an explicit claim in the text —
  `Ожидалось 2/2`, `решение верное (2/2)`, `по критериям 2/2`. Never from a
  bare `2/2` in a neighbouring cell: that is usually the bot's own score, and
  reading it as the expectation manufactures a disagreement.

The original short comment is kept as `comment_app` — the two are not always
the same claim, and a merge that overwrites evidence is not a merge.

Run with --dry-run first: it prints what it matched and changes nothing.
"""

from __future__ import annotations

import argparse
import json
import pathlib
import re
import sys
from typing import Any

ID_RE = re.compile(r"grade_[0-9a-f]{12,}", re.I)
PAIR_RE = re.compile(r"\b([0-4])\s*/\s*([0-4])\b")
# Only explicit statements of what the grade *should* have been. An earlier
# version also scanned short neighbouring cells for an "N/N" and picked up the
# bot's own score from the next column, inventing a mismatch that the tester
# never claimed.
EXPECT_PATTERNS = [
    re.compile(r"ожида\w*\s*[:\-–—]?\s*([0-4])\s*/\s*([0-4])", re.I),
    re.compile(r"(?:реш\w+|ответ\w*)\s+\w*\s*вер\w+\s*\(\s*([0-4])\s*/\s*([0-4])\s*\)", re.I),
    re.compile(r"по\s+критери\w+\s*[—\-–:]?\s*([0-4])\s*/\s*([0-4])", re.I),
    re.compile(r"должно\s+быть\s*([0-4])\s*/\s*([0-4])", re.I),
    re.compile(r"\bверн\w*\s*\(\s*([0-4])\s*/\s*([0-4])\s*\)", re.I),
]
WRONG_CUES = ("не прав", "неправ", "неверно", "ошиб", "занижен", "завышен")
OK_CUES = ("прав", "верно", "согласен", "корректно")


def cell_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def read_sheet(path: pathlib.Path, sheet: str | None) -> tuple[str, list[list[str]]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    names = wb.sheetnames
    if sheet and sheet in names:
        name = sheet
    elif sheet:
        # tolerate case and stray spaces in a hand-typed sheet name
        match = [n for n in names if n.strip().lower() == sheet.strip().lower()]
        name = match[0] if match else names[0]
        if not match:
            print(f"  sheet {sheet!r} not found; sheets are {names}, using {name!r}", file=sys.stderr)
    else:
        name = names[0]
    ws = wb[name]
    rows = [[cell_text(c) for c in row] for row in ws.iter_rows(values_only=True)]
    wb.close()
    return name, [r for r in rows if any(r)]


def find_id_column(rows: list[list[str]]) -> int | None:
    hits: dict[int, int] = {}
    for row in rows:
        for i, val in enumerate(row):
            if ID_RE.search(val):
                hits[i] = hits.get(i, 0) + 1
    return max(hits, key=hits.get) if hits else None  # type: ignore[arg-type]


def find_text_column(rows: list[list[str]], skip: set[int]) -> int | None:
    """The column carrying the most prose — the verdict text."""
    score: dict[int, int] = {}
    for row in rows:
        for i, val in enumerate(row):
            if i in skip or len(val) < 40:
                continue
            score[i] = score.get(i, 0) + len(val)
    return max(score, key=score.get) if score else None  # type: ignore[arg-type]


def find_verdict_column(rows: list[list[str]], skip: set[int]) -> int | None:
    score: dict[int, int] = {}
    for row in rows:
        for i, val in enumerate(row):
            if i in skip or not val or len(val) > 60:
                continue
            low = val.lower()
            if any(c in low for c in WRONG_CUES) or any(c in low for c in OK_CUES):
                score[i] = score.get(i, 0) + 1
    return max(score, key=score.get) if score else None  # type: ignore[arg-type]


def verdict_from(text: str) -> str:
    low = text.lower()
    # "НЕ прав" must beat "прав": check the negative cues first
    if any(c in low for c in WRONG_CUES):
        return "wrong"
    if any(c in low for c in OK_CUES):
        return "ok"
    return "unsure"


def expected_from(row: list[str], comment: str) -> tuple[int | None, int | None]:
    """What the tester said the grade should have been, or (None, None).

    Deliberately no fallback to a bare "N/N" in a neighbouring cell: that cell
    is usually the bot's score, and reading it as the expectation manufactures
    a disagreement out of an agreement.
    """
    for pattern in EXPECT_PATTERNS:
        if m := pattern.search(comment):
            return int(m.group(1)), int(m.group(2))
    return None, None


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", type=pathlib.Path)
    ap.add_argument("reports", nargs="?", default="collected-reports", type=pathlib.Path)
    ap.add_argument("--sheet", default="Прогоны")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    sheet_name, rows = read_sheet(args.xlsx, args.sheet)
    id_col = find_id_column(rows)
    if id_col is None:
        sys.exit(
            f"no column of grade_… ids on sheet {sheet_name!r}. "
            "Merging needs request_id to match rows to runs."
        )
    text_col = find_text_column(rows, {id_col})
    verdict_col = find_verdict_column(rows, {id_col, text_col} - {None})  # type: ignore[arg-type]
    print(f"sheet {sheet_name!r}: {len(rows)} rows | id col {id_col} | "
          f"text col {text_col} | verdict col {verdict_col}")

    by_id: dict[str, dict[str, Any]] = {}
    for row in rows:
        m = ID_RE.search(row[id_col] if id_col < len(row) else "")
        if not m:
            continue
        rid = m.group(0)
        comment = row[text_col] if text_col is not None and text_col < len(row) else ""
        vtext = row[verdict_col] if verdict_col is not None and verdict_col < len(row) else ""
        exp_b, exp_p = expected_from(row, comment)
        by_id[rid] = {
            "comment": comment,
            "verdict": verdict_from(vtext or comment),
            "expected_base": exp_b,
            "expected_presentation": exp_p,
        }
    print(f"matched {len(by_id)} rows carrying a request_id")

    files = {p.stem: p for p in args.reports.rglob("*.json")}
    hit = miss = added = replaced = 0
    for rid, info in by_id.items():
        path = files.get(rid)
        if path is None:
            miss += 1
            continue
        hit += 1
        data = json.loads(path.read_text(encoding="utf-8"))
        old = data.get("feedback") or {}
        if old.get("comment"):
            replaced += 1
        else:
            added += 1
        merged = {
            **old,
            **{k: v for k, v in info.items() if v not in (None, "")},
            "source": "xlsx",
        }
        # keep the in-app text: the two are not always the same claim
        if old.get("comment") and old.get("comment") != info["comment"]:
            merged["comment_app"] = old["comment"]
        merged.setdefault("tester", old.get("tester") or "Тестировщик")
        data["feedback"] = merged
        if not args.dry_run:
            path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"runs updated: {hit} (new feedback {added}, replaced {replaced}) | "
          f"ids not found among reports: {miss}")
    without = len(files) - hit
    print(f"runs still without a spreadsheet verdict: {without}")
    if args.dry_run:
        print("\n--dry-run: nothing written. Sample of what would be merged:")
        for rid, info in list(by_id.items())[:3]:
            print(f"  {rid}  [{info['verdict']}] "
                  f"ожидалось {info['expected_base']}/{info['expected_presentation']}")
            print(f"    {info['comment'][:160]}")


if __name__ == "__main__":
    main()
